import os
import re
import json
import time
import pandas as pd
from collections import OrderedDict
from dotenv import load_dotenv
from flask import Flask, render_template, request, jsonify, session
from werkzeug.utils import secure_filename
from prometheus_client import Counter, Histogram, generate_latest, CONTENT_TYPE_LATEST

# ==========================================
# 1. CONFIGURACIÓN INICIAL Y ENTORNO
# ==========================================

# Cargar variables del archivo .env
load_dotenv()

# Instanciar la aplicación Flask
app = Flask(__name__)
app.secret_key = os.getenv('FLASK_SECRET_KEY', 'clave_desarrollo_local_2026')

# Carpeta de destino para archivos subidos
UPLOAD_FOLDER = os.getenv('UPLOAD_FOLDER', 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

#POSTGRESSSQL
import psycopg2
from psycopg2.extras import RealDictCursor

# Configuración de conexión a PostgreSQL (TP10)
DATABASE_URL = os.getenv('DATABASE_URL', 'postgresql://devsecops_user:devsecops_password_2026@localhost:5432/naturgy_db')

def get_db_connection():
    try:
        import psycopg2
        from psycopg2.extras import RealDictCursor
        
        # Detecta si debe usar 'db' (Docker) o 'localhost' (ejecución directa)
        db_host = os.getenv('DB_HOST', 'localhost' if os.getenv('FLASK_ENV') == 'development' else 'db')
        
        conn = psycopg2.connect(
            host=db_host,
            port=os.getenv('DB_PORT', '5432'),
            dbname=os.getenv('POSTGRES_DB', 'naturgy_db'),
            user=os.getenv('POSTGRES_USER', 'postgres'),
            password=os.getenv('POSTGRES_PASSWORD', 'admin2026'),
            cursor_factory=RealDictCursor,
            connect_timeout=3
        )
        return conn
    except Exception as e:
        print(f"⚠️ Error conectando a PostgreSQL ({os.getenv('DB_HOST', 'localhost')}): {e}")
        return None

# ==========================================
# 2. MÉTRICAS DE OBSERVABILIDAD (PROMETHEUS)
# ==========================================

REQUEST_COUNT = Counter(
    'flask_http_request_total', 
    'Total de peticiones HTTP en la aplicacion', 
    ['method', 'endpoint', 'status_code']
)

REQUEST_LATENCY = Histogram(
    'flask_http_request_duration_seconds', 
    'Latencia de las peticiones HTTP en segundos', 
    ['endpoint']
)

FILE_UPLOAD_COUNT = Counter(
    'flask_file_uploads_total', 
    'Total de archivos subidos', 
    ['status']
)

@app.before_request
def start_timer():
    request._start_time = time.time()

@app.after_request
def record_metrics(response):
    if hasattr(request, '_start_time'):
        resp_time = time.time() - request._start_time
        endpoint = request.endpoint or 'unknown'
        REQUEST_LATENCY.labels(endpoint=endpoint).observe(resp_time)
        REQUEST_COUNT.labels(
            method=request.method, 
            endpoint=endpoint, 
            status_code=response.status_code
        ).inc()
    return response

# ==========================================
# 3. USUARIOS DESDE VARIABLES DE ENTORNO (.ENV)
# ==========================================
USUARIOS = {
    # --- LÍDERES ---
    "kparra": {
        "password": os.getenv('PASS_KPARRA') or os.getenv('PASS_LIDER_1'),
        "rol": "lider",
        "equipo": "TODOS"
    },
    "lherrera": {
        "password": os.getenv('PASS_LHERRERA') or os.getenv('PASS_LIDER_2'),
        "rol": "lider",
        "equipo": "TODOS"
    },

    # --- ASESORES ---
    "mquintana": {"password": os.getenv('PASS_MQUINTANA'), "rol": "asesor", "equipo": "RAC"},
    "aescobar":  {"password": os.getenv('PASS_AESCOBAR'),  "rol": "asesor", "equipo": "RAC"},
    "lganino":   {"password": os.getenv('PASS_LGANINO'),   "rol": "asesor", "equipo": "RAC"},
    "aherbas":   {"password": os.getenv('PASS_AHERBAS'),   "rol": "asesor", "equipo": "RAC"},
    "frios":     {"password": os.getenv('PASS_FRIOS'),     "rol": "asesor", "equipo": "RAC"},
    "carce":     {"password": os.getenv('PASS_CARCE'),     "rol": "asesor", "equipo": "RAC"},
    "arios":     {"password": os.getenv('PASS_ARIOS'),     "rol": "asesor", "equipo": "RAC"},
    "fruiz":     {"password": os.getenv('PASS_FRUIZ'),     "rol": "asesor", "equipo": "RAC"},
    "vgarcia":   {"password": os.getenv('PASS_VGARCIA'),   "rol": "asesor", "equipo": "RAC"},
    "nbogado":   {"password": os.getenv('PASS_NBOGADO'),   "rol": "asesor", "equipo": "RAC"},
    "vgonzalez": {"password": os.getenv('PASS_VGONZALEZ'), "rol": "asesor", "equipo": "RAC"},
    "sgonzalez": {"password": os.getenv('PASS_SGONZALEZ'), "rol": "asesor", "equipo": "RAC"},
}

# ==========================================
# 4. CARGA Y LECTURA DE NÓMINAS DESDE .ENV
# ==========================================

def cargar_nomina_env(var_name):
    try:
        raw = os.getenv(var_name, '{}').strip().strip("'\"")
        return json.loads(raw)
    except Exception as e:
        print(f"⚠️ Error al parsear {var_name} del .env: {e}")
        return {}

NOMINA_BAN = cargar_nomina_env('NOMINA_BAN')
NOMINA_SAN_JUAN = cargar_nomina_env('NOMINA_SAN_JUAN')
NOMINA_NOA = cargar_nomina_env('NOMINA_NOA')

@app.route('/get_nomina/<equipo>', methods=['GET'])
def get_nomina(equipo):
    """Devuelve la nómina del equipo incluyendo legajo, nombre, líder y turno"""
    equipo_upper = equipo.upper()
    nomina_lista = []

    # 1. Intentar consulta a PostgreSQL (en Docker)
    conn = get_db_connection()
    if conn:
        try:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT legajo, nombre, lider, turno FROM nominas WHERE UPPER(equipo) LIKE %s;",
                    (f"%{equipo_upper}%",)
                )
                filas = cur.fetchall()
                for f in filas:
                    nomina_lista.append({
                        "legajo": str(f.get('legajo', '---')),
                        "nombre": f.get('nombre', '---'),
                        "lider": f.get('lider', '---'),
                        "turno": f.get('turno', '---')
                    })
            conn.close()
        except Exception as e:
            print(f"⚠️ Error consultando BD de nómina: {e}")

    # 2. Respaldo desde el .env (local sin Docker)
    if not nomina_lista:
        if "BAN" in equipo_upper:
            dict_env = NOMINA_BAN
        elif "SAN JUAN" in equipo_upper:
            dict_env = NOMINA_SAN_JUAN
        elif "NOA" in equipo_upper:
            dict_env = NOMINA_NOA
        else:
            dict_env = {}

        for legajo_key, info in dict_env.items():
            if isinstance(info, dict):
                nomina_lista.append({
                    "legajo": str(legajo_key),
                    "nombre": info.get("nombre", "---"),
                    "lider": info.get("lider", "---"),
                    "turno": info.get("turno", "---")
                })

    if nomina_lista:
        return jsonify({"success": True, "data": nomina_lista})

    return jsonify({"success": False, "message": f"Nómina no encontrada para {equipo}"})

# ==========================================
# 5. FUNCIONES AUXILIARES DE PROCESAMIENTO
# ==========================================

def clean_str(val):
    """Limpia valores nulos o vacíos convirtiéndolos a string plano."""
    if pd.isna(val) or val is None:
        return ''
    return str(val).strip()


def procesar_archivo_reporte(filepath):
    """
    Lee planillas individuales o grupales de logueo (CSV/Excel)
    y las transforma en DataFrames de 'Hora de Entrada' y 'Hora de Salida'.
    """
    split_res = os.path.splitext(filepath)
    ext = str(split_res[1]).lower() if isinstance(split_res, (list, tuple)) else str(split_res).lower()
    df_raw = None
    
    # 1. Intentar lectura como Excel (.xlsx, .xls)
    if ext in ['.xlsx', '.xls']:
        for engine in ['openpyxl', 'xlrd', None]:
            try:
                df_raw = pd.read_excel(filepath, sheet_name=0, header=None, engine=engine)
                if df_raw is not None and not df_raw.empty:
                    break
            except Exception:
                pass

    # 2. Intentar lectura como CSV
    if df_raw is None or df_raw.empty:
        for enc in ['utf-8', 'latin1', 'cp1252']:
            for sep in [';', ',', '\t', None]:
                try:
                    df_raw = pd.read_csv(filepath, sep=sep, encoding=enc, header=None, on_bad_lines='skip', engine='python')
                    if df_raw is not None and not df_raw.empty and df_raw.shape[1] > 1:
                        break
                except Exception:
                    pass
            if df_raw is not None and not df_raw.empty and df_raw.shape[1] > 1:
                break

    if df_raw is None or df_raw.empty:
        return []

    raw_text_list = [clean_str(x) for x in df_raw.values.flatten()]
    raw_text = ' '.join([x for x in raw_text_list if x])

    # CASO A: Reporte Individual por Agente
    if 'AGENTE:' in raw_text.upper() or 'REPORTE MENSUAL POR AGENTE' in raw_text.upper():
        agente = 'ASESOR'
        for r_idx in range(min(10, len(df_raw))):
            row_vals = [clean_str(x) for x in df_raw.iloc[r_idx].values]
            row_str = ' '.join([x for x in row_vals if x])
            if 'AGENTE:' in row_str.upper():
                m = re.search(r'AGENTE:\s*([^;|\n\r]+)', row_str, re.IGNORECASE)
                if m:
                    agente = m.group(1).strip()
                    break

        header_idx = None
        for r_idx in range(min(15, len(df_raw))):
            row_vals = [clean_str(x) for x in df_raw.iloc[r_idx].values]
            row_str = ' '.join([x for x in row_vals if x]).upper()
            if 'DESCRIPCION' in row_str or 'DESCRIPCIÓN' in row_str:
                header_idx = r_idx
                break

        if header_idx is not None:
            header_row = [clean_str(x) for x in df_raw.iloc[header_idx].values]
            df_data = df_raw.iloc[header_idx + 1:].copy()
            df_data.columns = header_row
        else:
            df_data = df_raw.copy()

        desc_col = None
        for col in df_data.columns:
            if 'DESCRIP' in str(col).upper():
                desc_col = col
                break

        if desc_col is None:
            return []

        df_in = df_data[df_data[desc_col].astype(str).str.strip().str.upper() == 'HORA DE ENTRADA']
        df_out = df_data[df_data[desc_col].astype(str).str.strip().str.upper() == 'HORA DE SALIDA']

        if df_in.empty or df_out.empty:
            return []

        ignore_keywords = ['GRUPO', 'DESCRIP', 'TOTAL', 'UNNAMED', 'NAN', 'NONE']
        date_cols = []
        for col in df_data.columns:
            c_str = str(col).strip()
            c_upper = c_str.upper()
            if c_str != '' and not any(k in c_upper for k in ignore_keywords):
                date_cols.append(col)

        row1 = {'Asesor': agente, 'DESCRIPCION': 'Hora de Entrada'}
        row2 = {'Asesor': '', 'DESCRIPCION': 'Hora de Salida'}

        for col in date_cols:
            col_name = str(col).strip()
            v_in_vals = df_in[col].values
            v_out_vals = df_out[col].values
            
            v_in = v_in_vals[0] if len(v_in_vals) > 0 else '---'
            v_out = v_out_vals[0] if len(v_out_vals) > 0 else '---'
            
            str_in = clean_str(v_in)
            str_out = clean_str(v_out)
            
            row1[col_name] = str_in if str_in and str_in.lower() != 'nan' else '---'
            row2[col_name] = str_out if str_out and str_out.lower() != 'nan' else '---'

        return [pd.DataFrame([row1, row2])]

    # CASO B: Reporte Grupal
    blocks = []
    current_agent = None
    dates = []

    for r_idx in range(len(df_raw)):
        row_vals = [clean_str(x) for x in df_raw.iloc[r_idx].values]
        row_str = ' '.join([x for x in row_vals if x]).upper()
        
        v0 = row_vals[0] if len(row_vals) > 0 else ''
        v1 = row_vals[1] if len(row_vals) > 1 else ''

        if v1.upper() in ['DESCRIPCION', 'DESCRIPCIÓN']:
            current_agent = v0 if v0 else 'ASESOR'
            if r_idx + 1 < len(df_raw):
                next_vals = [clean_str(x) for x in df_raw.iloc[r_idx + 1].values]
                dates = [v for v in next_vals if v and any(t in v.lower() for t in ['oct', 'sep', '/', '-'])]
            continue

        if 'HORA DE ENTRADA' in row_str:
            in_vals = [v for v in row_vals if v]
            out_vals = []
            if r_idx + 1 < len(df_raw):
                out_vals = [clean_str(x) for x in df_raw.iloc[r_idx + 1].values if clean_str(x)]

            row1 = {'Asesor': current_agent or 'ASESOR', 'DESCRIPCION': 'Hora de Entrada'}
            row2 = {'Asesor': '', 'DESCRIPCION': 'Hora de Salida'}

            in_times = [v for v in in_vals if v not in [current_agent, 'DESCRIPCION', 'Hora de Entrada'] and v.lower() != 'nan']
            out_times = [v for v in out_vals if v not in ['DESCRIPCION', 'Hora de Salida'] and v.lower() != 'nan']

            for d_idx, d_name in enumerate(dates):
                v_in = in_times[d_idx] if d_idx < len(in_times) else '---'
                v_out = out_times[d_idx] if d_idx < len(out_times) else '---'
                row1[d_name] = v_in if v_in and v_in.lower() != 'nan' else '---'
                row2[d_name] = v_out if v_out and v_out.lower() != 'nan' else '---'

            blocks.append(pd.DataFrame([row1, row2]))

    return blocks


# ==========================================
# 6. RUTAS PRINCIPALES Y AUTENTICACIÓN
# ==========================================

@app.route('/metrics', methods=['GET'])
def metrics():
    """Expone métricas en formato compatible con Prometheus"""
    return generate_latest(), 200, {'Content-Type': CONTENT_TYPE_LATEST}


@app.route('/')
def index():
    """Renderea la vista web principal"""
    return render_template('index.html')

@app.route('/login', methods=['POST'])
def login():
    """Valida usuarios consultando PostgreSQL o el fallback de .env"""
    data = request.get_json() or {}
    usuario_input = data.get('usuario', '').strip().lower()
    password_input = data.get('password', '').strip()
    equipo_solicitado = data.get('equipo', '').strip()

    # 1. INTENTAR VALIDAR CONTRA POSTGRESQL
    conn = get_db_connection()
    if conn:
        try:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT usuario, password, rol, equipo FROM usuarios WHERE LOWER(usuario) = %s;",
                    (usuario_input,)
                )
                user_db = cur.fetchone()
            conn.close()

            if user_db and user_db['password'] == password_input:
                if user_db['rol'] != 'lider' and user_db['equipo'] != "TODOS" and user_db['equipo'] != equipo_solicitado:
                    return jsonify({
                        "success": False, 
                        "message": f"Acceso denegado: El usuario pertenece al equipo '{user_db['equipo']}'."
                    }), 403

                session['usuario'] = user_db['usuario']
                session['rol'] = user_db['rol']
                session['equipo'] = equipo_solicitado if user_db['rol'] == 'lider' else user_db['equipo']

                return jsonify({
                    "success": True,
                    "usuario": user_db['usuario'],
                    "rol": user_db['rol'],
                    "equipo": session['equipo']
                })
        except Exception as e:
            print(f"⚠️ Fallo consulta DB, usando fallback local: {e}")

    # 2. FALLBACK A USUARIOS LOCALES (.ENV)
    pass_kparra = os.getenv('PASS_KPARRA', 'admin2026')
    pass_lherrera = os.getenv('PASS_LHERRERA', 'admin2026')

    if usuario_input in ['kparra', 'lherrera']:
        pass_valida = pass_kparra if usuario_input == 'kparra' else pass_lherrera
        if password_input == pass_valida:
            session['usuario'] = usuario_input
            session['rol'] = 'lider'
            session['equipo'] = equipo_solicitado
            return jsonify({
                "success": True,
                "usuario": usuario_input,
                "rol": "lider",
                "equipo": equipo_solicitado
            })

    return jsonify({"success": False, "message": "Usuario o contraseña incorrectos"}), 401


@app.route('/logout', methods=['POST'])
def logout():
    """Cierra la sesión del usuario"""
    session.clear()
    return jsonify({"success": True})

# ==========================================
# 7. CARGA Y CONSOLIDACIÓN DE REPORTES
# ==========================================

@app.route('/upload', methods=['POST'])
def upload_file():
    """Permite al Líder subir planillas de repocfrtes (.csv, .xlsx, .xls)"""
    if session.get('rol') != 'lider':
        return jsonify({"success": False, "message": "Se requieren permisos de Líder."}), 403

    if 'file' not in request.files:
        return jsonify({"success": False, "message": "No se adjuntó ningún archivo."}), 400

    file = request.files['file']
    if file.filename == '' or file.filename == '[object FileList]':
        return jsonify({"success": False, "message": "Por favor selecciona un archivo válido."}), 400

    filename = secure_filename(file.filename)
    ext = os.path.splitext(filename)[1].lower()
    
    if ext not in ['.csv', '.xlsx', '.xls']:
        return jsonify({"success": False, "message": "Formato no permitido. Debe ser .csv, .xlsx o .xls."}), 400

    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(filepath)
    FILE_UPLOAD_COUNT.labels(status='success').inc()

    return jsonify({
        "success": True,
        "message": f"Reporte '{filename}' cargado correctamente."
    })

@app.route('/obtener_reporte', methods=['GET'])
def obtener_reporte():
    """Consolida todos los archivos en uploads/ y genera la tabla HTML"""
    if 'usuario' not in session:
        return jsonify({"success": False, "message": "Sesión no iniciada"}), 401

    equipo_activo = session.get('equipo')

    archivos = [
        os.path.join(app.config['UPLOAD_FOLDER'], f) 
        for f in os.listdir(app.config['UPLOAD_FOLDER']) 
        if f.endswith(('.csv', '.xlsx', '.xls'))
    ]

    if not archivos:
        return jsonify({
            "success": False, 
            "message": "No hay reportes cargados aún para este equipo."
        })

    try:
        all_blocks = []
        for archivo in archivos:
            blocks = procesar_archivo_reporte(archivo)
            all_blocks.extend(blocks)

        if not all_blocks:
            return jsonify({"success": False, "message": "No se pudieron procesar los datos de las planillas."})

        seen_advisors = set()
        dedup_blocks = []
        for block in reversed(all_blocks):
            adv_name = clean_str(block['Asesor'].iloc[0]) if len(block) > 0 and 'Asesor' in block.columns else ''
            if adv_name and adv_name != 'ASESOR':
                if adv_name not in seen_advisors:
                    seen_advisors.add(adv_name)
                    dedup_blocks.append(block)
            else:
                dedup_blocks.append(block)

        dedup_blocks.reverse()
        df_consolidado = pd.concat(dedup_blocks, ignore_index=True).fillna('---')

        html_tabla = df_consolidado.to_html(
            classes='table table-striped tabla-reporte', 
            index=False
        )

        return jsonify({
            "success": True,
            "equipo": equipo_activo,
            "html_tabla": html_tabla
        })

    except Exception as e:
        return jsonify({
            "success": False, 
            "message": f"Error al procesar la consolidación: {str(e)}"
        }), 500


# ==========================================
# 8. GESTIÓN DE CALIDAD Y MONITOREOS (RAC)
# ==========================================

@app.route('/upload_calidad_resumen', methods=['POST'])
def upload_calidad_resumen():
    """Procesa planillas de Calidad para construir los resúmenes por RAC y TL"""
    if session.get('rol') and session.get('rol') != 'lider':
        return jsonify({"success": False, "message": "Se requieren permisos de Líder."}), 403

    if 'file' not in request.files:
        return jsonify({"success": False, "message": "No se adjuntó ningún archivo."}), 400

    file = request.files['file']
    if file.filename == '' or file.filename == '[object FileList]':
        return jsonify({"success": False, "message": "Archivo no seleccionado o no válido."}), 400

    try:
        upload_dir = app.config.get('UPLOAD_FOLDER', 'uploads')
        filename = secure_filename(file.filename)
        filepath = os.path.join(upload_dir, 'resumen_' + filename)
        file.save(filepath)

        ext = os.path.splitext(filepath)[1].lower()
        df_raw = None

        if ext in ['.xlsx', '.xls', '.xlsm']:
            try:
                xl = pd.ExcelFile(filepath)
                target_sheet = next((s for s in xl.sheet_names if 'CONSOLIDADOS POR RAC' in s.upper() or 'CONSOLIDADO' in s.upper()), xl.sheet_names[0] if xl.sheet_names else None)

                if target_sheet:
                    for engine in ['openpyxl', 'xlrd', None]:
                        try:
                            df_raw = pd.read_excel(filepath, sheet_name=target_sheet, header=None, engine=engine)
                            if df_raw is not None and not df_raw.empty:
                                break
                        except Exception:
                            continue
            except Exception as e_xl:
                print("Aviso al abrir libro Excel:", str(e_xl))

        if df_raw is None or df_raw.empty:
            return jsonify({"success": False, "message": "No se encontraron datos válidos en la hoja seleccionada."}), 400

        header_idx = None
        for r_idx in range(len(df_raw)):
            row_str = ' '.join([clean_str(x).upper() for x in df_raw.iloc[r_idx].values if pd.notna(x)])
            if 'LEGAJO' in row_str and 'ASESOR' in row_str and ('LÍDER' in row_str or 'LIDER' in row_str):
                header_idx = r_idx
                break

        if header_idx is None:
            return jsonify({"success": False, "message": "No se encontró la tabla principal en la planilla de calidad."}), 400

        raw_headers = [clean_str(x).upper() for x in df_raw.iloc[header_idx].values]
        
        idx_map = {}
        for c_idx, h_text in enumerate(raw_headers):
            if ('LÍDER' in h_text or 'LIDER' in h_text) and 'Líder' not in idx_map:
                idx_map['Líder'] = c_idx
            elif 'ASESOR' in h_text and 'Asesor' not in idx_map:
                idx_map['Asesor'] = c_idx
            elif '1' in h_text and 'MONITOREO' in h_text and '1º Monitoreo' not in idx_map:
                idx_map['1º Monitoreo'] = c_idx
            elif '2' in h_text and 'MONITOREO' in h_text and '2º Monitoreo' not in idx_map:
                idx_map['2º Monitoreo'] = c_idx
            elif '3' in h_text and 'MONITOREO' in h_text and '3º Monitoreo' not in idx_map:
                idx_map['3º Monitoreo'] = c_idx
            elif 'QM' in h_text and 'QM' not in idx_map:
                idx_map['QM'] = c_idx
            elif ('NOTA' in h_text or 'FINAL' in h_text) and 'Nota Final' not in idx_map:
                idx_map['Nota Final'] = c_idx

        def formatear_nota(num):
            try:
                val_num = float(num)
                if val_num >= 85.0:
                    return f'<span class="badge-cumple">{val_num:.2f}%</span>'
                else:
                    return f'<span class="badge-nocumple">{val_num:.2f}%</span>'
            except Exception:
                return '---'

        def parse_float(val):
            v_str = clean_str(val)
            if not v_str or v_str.upper() in ['', 'NAN', 'NONE', '---', '#¡DIV/0!']:
                return None
            try:
                num = float(v_str.replace('%', '').replace(',', '.').strip())
                if 0 < num <= 1.0:
                    num = num * 100
                return num
            except Exception:
                return None

        def parse_int(val):
            v_str = clean_str(val)
            if not v_str or v_str.upper() in ['', 'NAN', 'NONE', '---', '#¡DIV/0!']:
                return 0
            try:
                return int(float(v_str.replace(',', '.').strip()))
            except Exception:
                return 0

        rows_asesores = []
        lideres_definidos = ['HERRERA FERNANDEZ, LEILA SABRINA', 'PARRA YAÑEZ, KARINA BEATRIZ']
        matrix_data = df_raw.iloc[header_idx + 1:].values

        for row in matrix_data:
            lider_val = clean_str(row[idx_map['Líder']]) if 'Líder' in idx_map and idx_map['Líder'] < len(row) else ''
            asesor_val = clean_str(row[idx_map['Asesor']]) if 'Asesor' in idx_map and idx_map['Asesor'] < len(row) else ''

            if asesor_val.upper() in ['', 'ASESOR', 'NAN', 'NONE', 'NONETYPE']:
                continue

            lider_coincide = next((l_target for l_target in lideres_definidos if l_target.upper() in lider_val.upper()), None)

            if lider_coincide:
                m1_num = parse_float(row[idx_map['1º Monitoreo']]) if '1º Monitoreo' in idx_map else None
                m2_num = parse_float(row[idx_map['2º Monitoreo']]) if '2º Monitoreo' in idx_map else None
                m3_num = parse_float(row[idx_map['3º Monitoreo']]) if '3º Monitoreo' in idx_map else None
                qm_num = parse_int(row[idx_map['QM']]) if 'QM' in idx_map else 0
                nota_num = parse_float(row[idx_map['Nota Final']]) if 'Nota Final' in idx_map else None

                rows_asesores.append({
                    'Líder': lider_coincide,
                    'Asesor': asesor_val,
                    '1º Monitoreo': formatear_nota(m1_num) if m1_num is not None else '---',
                    '2º Monitoreo': formatear_nota(m2_num) if m2_num is not None else '---',
                    '3º Monitoreo': formatear_nota(m3_num) if m3_num is not None else '---',
                    'QM': str(qm_num),
                    'Nota Final': formatear_nota(nota_num) if nota_num is not None else '---',
                    '_qm_raw': qm_num,
                    '_nota_raw': nota_num
                })

        rows_tl = []
        total_campana_qm = 0
        notas_campana = []

        for l_target in lideres_definidos:
            asesores_lider = [r for r in rows_asesores if r['Líder'] == l_target]
            sum_qm = sum(r['_qm_raw'] for r in asesores_lider)
            notas_validas = [r['_nota_raw'] for r in asesores_lider if r['_nota_raw'] is not None]
            prom_nota = (sum(notas_validas) / len(notas_validas)) if notas_validas else None

            total_campana_qm += sum_qm
            notas_campana.extend(notas_validas)

            rows_tl.append({
                'Team Leader': l_target,
                'Monitoreos': str(sum_qm),
                'TOTAL': formatear_nota(prom_nota) if prom_nota is not None else '---'
            })

        prom_campana = (sum(notas_campana) / len(notas_campana)) if notas_campana else None
        rows_tl.append({
            'Team Leader': 'TOTAL CAMPAÑA',
            'Monitoreos': str(total_campana_qm),
            'TOTAL': formatear_nota(prom_campana) if prom_campana is not None else '---'
        })

        for r in rows_asesores:
            r.pop('_qm_raw', None)
            r.pop('_nota_raw', None)

        cols_finales_asesor = ['Líder', 'Asesor', '1º Monitoreo', '2º Monitoreo', '3º Monitoreo', 'QM', 'Nota Final']
        df_asesores = pd.DataFrame(rows_asesores, columns=cols_finales_asesor) if rows_asesores else pd.DataFrame(columns=cols_finales_asesor)

        cols_finales_tl = ['Team Leader', 'Monitoreos', 'TOTAL']
        df_tl = pd.DataFrame(rows_tl, columns=cols_finales_tl)

        html_asesores = df_asesores.fillna('---').to_html(classes='table table-striped tabla-reporte', index=False, escape=False) if not df_asesores.empty else '<p class="text-muted">No se encontraron registros para las líderes seleccionadas.</p>'
        html_tl = df_tl.to_html(classes='table table-bordered tabla-tl-custom', index=False, escape=False)
        
        html_tl_container = f'''
        <style>
            .tabla-tl-custom {{ border: 2px solid #000 !important; text-align: center; font-weight: bold; }}
            .tabla-tl-custom th {{ background-color: #F79646 !important; color: white !important; text-align: center !important; border: 1px solid #000 !important; }}
            .tabla-tl-custom td {{ border: 1px solid #000 !important; vertical-align: middle !important; }}
        </style>
        <div style="display: flex; justify-content: flex-end; margin-bottom: 25px;">
            <div style="min-width: 420px;">{html_tl}</div>
        </div>
        '''

        return jsonify({
            "success": True,
            "message": "Planilla Resumen procesada correctamente.",
            "html_tabla": html_tl_container + html_asesores
        })

    except Exception as e:
        return jsonify({"success": False, "message": f"Error: {type(e).__name__} - {str(e)}"}), 400


# ==========================================
# 9. REGISTRO Y MATRIZ DE LOGUEO
# ==========================================

@app.route('/upload_logueo', methods=['POST'])
def upload_logueo():
    """Procesa e integra planillas de fichada diaria de logueo con selectores de estados"""
    if 'file' not in request.files:
        return jsonify({"success": False, "message": "No se adjuntó ningún archivo."}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({"success": False, "message": "Archivo no válido."}), 400

    try:
        upload_dir = app.config.get('UPLOAD_FOLDER', 'uploads')
        original_filename = secure_filename(file.filename)
        filepath = os.path.join(upload_dir, 'logueo_' + original_filename)
        file.save(filepath)

        ext = os.path.splitext(filepath)[1].lower()
        asesor_nombre = request.form.get('asesor_nombre', '').strip()
        df_raw = None

        if ext in ['.xlsx', '.xlsm']:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(filepath, data_only=True)
                ws = wb.active
                for rng in list(ws.merged_cells.ranges):
                    top_left_val = ws.cell(rng.min_row, rng.min_col).value
                    ws.unmerge_cells(str(rng))
                    for r_idx in range(rng.min_row, rng.max_row + 1):
                        for c_idx in range(rng.min_col, rng.max_col + 1):
                            ws.cell(r_idx, c_idx).value = top_left_val

                data_matrix = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
                temp_df = pd.DataFrame(data_matrix)
                if temp_df is not None and not temp_df.empty:
                    df_raw = temp_df
            except Exception:
                pass

        if df_raw is None or df_raw.empty:
            for enc in ['utf-8', 'latin1', 'cp1252']:
                try:
                    df_raw = pd.read_csv(filepath, sep=None, encoding=enc, header=None, engine='python')
                    if df_raw is not None and not df_raw.empty:
                        break
                except Exception:
                    continue

        if df_raw is None or df_raw.empty:
            return jsonify({"success": False, "message": "No se pudieron extraer datos del archivo subido."}), 400

        df_raw = df_raw.ffill(axis=1)

        if not asesor_nombre:
            for r in range(min(20, len(df_raw))):
                for c in range(min(20, df_raw.shape[1])):
                    val = clean_str(df_raw.iloc[r, c]).replace('\xa0', ' ').strip()
                    if 'AGENTE:' in val.upper():
                        partes = val.split('AGENTE:', 1) if 'AGENTE:' in val else val.split('Agente:', 1)
                        if len(partes) > 1 and partes[1].strip():
                            asesor_nombre = partes[1].split('|')[0].strip().upper()
                            break
                if asesor_nombre:
                    break

        if not asesor_nombre:
            asesor_nombre = os.path.splitext(original_filename)[0].replace('logueo_', '').replace('_', ' ').strip().upper()

        idx_entrada = next((r for r in range(len(df_raw)) if 'ENTRADA' in ' '.join([clean_str(x).upper() for x in df_raw.iloc[r].values])), None)
        idx_salida = next((r for r in range(len(df_raw)) if 'SALIDA' in ' '.join([clean_str(x).upper() for x in df_raw.iloc[r].values])), None)

        if idx_entrada is None or idx_salida is None:
            return jsonify({"success": False, "message": "No se detectaron filas de fichada."}), 400

        row_fechas = df_raw.iloc[max(0, idx_entrada - 1)]
        row_entrada = df_raw.iloc[idx_entrada]
        row_salida = df_raw.iloc[idx_salida]

        logueo_entrada = OrderedDict({"Asesor": asesor_nombre, "Tipo": "Entrada"})
        logueo_salida = OrderedDict({"Asesor": asesor_nombre, "Tipo": "Salida"})

        for c_idx in range(2, min(34, df_raw.shape[1])):
            val_f = clean_str(row_fechas[c_idx]) if c_idx < len(row_fechas) else ''
            if any(t in val_f.upper() for t in ['TOTAL', 'TOTALES', 'SUMA', 'Q1', 'Q2']):
                continue

            fecha_label = val_f if val_f and val_f.upper() not in ['NAN', 'NONE'] else f"Día {c_idx - 1:02d}"
            h_in = clean_str(row_entrada[c_idx]) if c_idx < len(row_entrada) else ''
            h_out = clean_str(row_salida[c_idx]) if c_idx < len(row_salida) else ''

            logueo_entrada[fecha_label] = h_in if h_in and h_in.upper() not in ['NAN', 'NONE', '---'] else '---'
            logueo_salida[fecha_label] = h_out if h_out and h_out.upper() not in ['NAN', 'NONE', '---'] else '---'

        asesor_key = " ".join(asesor_nombre.upper().split())
        if 'logueos_acumulados' not in session:
            session['logueos_acumulados'] = {}

        session['logueos_acumulados'][f"{asesor_key}_ENTRADA"] = dict(logueo_entrada)
        session['logueos_acumulados'][f"{asesor_key}_SALIDA"] = dict(logueo_salida)

        todas_las_columnas = set()
        for reg in session['logueos_acumulados'].values():
            todas_las_columnas.update(reg.keys())

        cols_fijas = ["Asesor", "Tipo"]
        cols_fechas = [c for c in todas_las_columnas if c not in cols_fijas]

        def criterio_orden_fechas(col_nombre):
            texto = str(col_nombre).strip()
            match_fecha = re.search(r'(\d{1,2})[/.-](\d{1,2})(?:[/.-](\d{2,4}))?', texto)
            if match_fecha:
                dia, mes = int(match_fecha.group(1)), int(match_fecha.group(2))
                anio = int(match_fecha.group(3)) if match_fecha.group(3) else 2026
                return (1, anio, mes, dia)
            match_dia = re.search(r'\d+', texto)
            if match_dia:
                return (2, 2026, 1, int(match_dia.group(0)))
            return (3, 9999, 12, 31)

        master_cols = cols_fijas + sorted(cols_fechas, key=criterio_orden_fechas)
        session['logueos_master_cols'] = master_cols
        session.modified = True

        registros = list(session['logueos_acumulados'].values())
        asesores_map = OrderedDict()
        for reg in registros:
            nombre = reg.get("Asesor", "")
            asesores_map.setdefault(nombre, []).append(reg)

        filas_consolidadas = []
        lista_asesores = list(asesores_map.keys())

        for idx, nombre in enumerate(lista_asesores):
            filas_consolidadas.extend(asesores_map[nombre])
            if idx < len(lista_asesores) - 1:
                filas_consolidadas.append({col: '' for col in master_cols})

        df_logueos = pd.DataFrame(filas_consolidadas)
        df_logueos = df_logueos[[c for c in master_cols if c in df_logueos.columns]]

        df_logueos['Asesor_Display'] = df_logueos['Asesor']
        df_logueos.loc[df_logueos['Asesor_Display'] == df_logueos['Asesor_Display'].shift(1), 'Asesor_Display'] = ''
        df_logueos['Asesor'] = df_logueos['Asesor_Display']
        df_logueos = df_logueos.drop(columns=['Asesor_Display'])

        FERIADOS_DDMM = {
            "01/01": "Año Nuevo", "16/02": "Carnaval", "17/02": "Carnaval", "24/03": "Memoria",
            "02/04": "Malvinas", "01/05": "Trabajador", "25/05": "Revolución de Mayo",
            "17/06": "Güemes", "20/06": "Bandera", "09/07": "Independencia", "17/08": "San Martín",
            "12/10": "Diversidad Cultural", "20/11": "Soberanía Nacional", "08/12": "Inmaculada Concepción", "25/12": "Navidad"
        }

        opciones_letras = ['F', 'FER', 'V', 'RH', 'E', 'AI', 'EX', 'L', 'ART', 'C', 'B']

        for col in [c for c in df_logueos.columns if c not in ['Asesor', 'Tipo']]:
            texto_col = str(col).strip()
            es_feriado_col = False
            match = re.search(r'(\d{1,2})[/.-](\d{1,2})', texto_col)
            if match:
                p1, p2 = match.group(1).zfill(2), match.group(2).zfill(2)
                if f"{p1}/{p2}" in FERIADOS_DDMM or f"{p2}/{p1}" in FERIADOS_DDMM:
                    es_feriado_col = True

            def procesar_celda(valor, es_fer=es_feriado_col):
                val_str = str(valor).strip().upper() if valor is not None and str(valor).strip() != 'nan' else ''
                if val_str == '':
                    return ''
                if ':' in val_str:
                    return val_str

                val_actual = 'FER' if es_fer else 'F' if val_str in ['---', 'NONE', 'NAN', 'F', 'FER'] else val_str
                opts_html = [f'<option value="{opt}" {"selected" if opt == val_actual else ""}>{opt}</option>' for opt in opciones_letras]
                return f'<select class="select-estado-logueo estado-{val_actual}" onchange="cambiarColorEstado(this)">{"".join(opts_html)}</select>'

            df_logueos[col] = df_logueos[col].apply(lambda x, f=es_feriado_col: procesar_celda(x, f))

        html_tabla = df_logueos.fillna('').to_html(classes='table table-bordered table-striped tabla-logueos', index=False, escape=False)

        return jsonify({
            "success": True,
            "message": f"Logueo procesado correctamente para '{asesor_key}'.",
            "html_tabla": html_tabla
        })      

    except Exception as e:
        return jsonify({"success": False, "message": f"Error: {type(e).__name__} - {str(e)}"}), 500


@app.route('/limpiar_logueos', methods=['POST'])
def limpiar_logueos():
    """Limpia los logueos acumulados en la sesión actual"""
    session.pop('logueos_acumulados', None)
    session.pop('logueos_master_cols', None)
    session.modified = True
    return jsonify({"success": True, "message": "Se han eliminado todos los reportes acumulados."})


# ==========================================
# 10. ARRANQUE DEL SERVIDOR FLASK
# ==========================================

if __name__ == '__main__':
    port = int(os.getenv('PORT', 5000))
    print(f"\n🚀 Servidor activo en http://127.0.0.1:{port}\n")
    app.run(debug=True, host='127.0.0.1', port=port, use_reloader=False)